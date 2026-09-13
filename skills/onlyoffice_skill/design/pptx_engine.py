"""
Motor de diseño premium para presentaciones PowerPoint (.pptx).
Proporciona creación y estructuración de diapositivas ejecutivas con layouts modulares:
title, agenda, content, two_columns, three_columns, kpi_metrics, conclusion, y tablas nativas.
"""

import logging
from typing import Optional, List, Dict, Any

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    PPTX_AVAILABLE = True
except Exception as e:
    PPTX_AVAILABLE = False

from skills.onlyoffice_skill.design.themes import (
    get_theme,
    hex_to_rgb,
    lighten_color,
    darken_color,
)

logger = logging.getLogger(__name__)

SUPPORTED_PPTX_LAYOUTS = [
    "title",
    "agenda",
    "content",
    "two_columns",
    "three_columns",
    "kpi_metrics",
    "conclusion",
]


def _ensure_widescreen(prs):
    """Configura la presentación en relación panorámica moderna 16:9 si aún no lo está."""
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)


def _add_header(slide, title: str, theme: Dict[str, Any], subtitle: Optional[str] = None):
    """Añade un encabezado corporativo estándar con barra divisoria de acento a la diapositiva."""
    p_rgb = RGBColor(*hex_to_rgb(theme["primary"]))
    s_rgb = RGBColor(*hex_to_rgb(theme["secondary"]))
    m_rgb = RGBColor(*hex_to_rgb(theme["text_muted"]))
    font_name = theme["font_heading"]

    # Caja de título superior
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)

    p = tf.paragraphs[0]
    p.text = title
    p.font.name = font_name
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = p_rgb

    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = font_name
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = m_rgb
        p_sub.space_before = Pt(2)

    # Barra decorativa de acento horizontal
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8),
        Inches(1.5),
        Inches(2.5),
        Inches(0.05)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = s_rgb
    accent_bar.line.fill.background()  # Sin borde


def add_slide(
    prs,
    layout_type: str = "content",
    data: Optional[Dict[str, Any]] = None,
    theme_name: str = "corporate_blue",
    speaker_notes: Optional[str] = None,
):
    """
    Crea y añade una diapositiva profesional con diseño de alta calidad según el layout elegido.
    layout_type: 'title', 'agenda', 'content', 'two_columns', 'three_columns', 'kpi_metrics', 'conclusion'
    """
    if not PPTX_AVAILABLE:
        return None

    _ensure_widescreen(prs)
    theme = get_theme(theme_name)
    data = data or {}
    font_heading = theme["font_heading"]
    font_body = theme["font_body"]

    p_rgb = RGBColor(*hex_to_rgb(theme["primary"]))
    s_rgb = RGBColor(*hex_to_rgb(theme["secondary"]))
    bg_light_rgb = RGBColor(*hex_to_rgb(theme["bg_light"]))
    border_rgb = RGBColor(*hex_to_rgb(theme["border"]))
    txt_rgb = RGBColor(*hex_to_rgb(theme["text_primary"]))
    m_rgb = RGBColor(*hex_to_rgb(theme["text_muted"]))

    # Usamos layout en blanco (índice 6) para control de precisión tipográfica y cajas
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank_layout)

    norm_layout = layout_type.lower().strip()

    if norm_layout == "title":
        # Portada Ejecutiva de Presentación
        # Fondo temático sólido o bloque elegante
        use_dark = data.get("dark_mode", True)
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = p_rgb if use_dark else bg_light_rgb
        bg_shape.line.fill.background()

        # Banda de acento lateral
        side_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.8)
        )
        side_bar.fill.solid()
        side_bar.fill.fore_color.rgb = s_rgb
        side_bar.line.fill.background()

        # Caja de título y subtítulo
        title_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.5), Inches(3.5))
        tf = title_box.text_frame
        tf.word_wrap = True

        p_cat = tf.paragraphs[0]
        p_cat.text = (data.get("category") or "PRESENTACIÓN CORPORATIVA").upper()
        p_cat.font.name = font_heading
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = s_rgb if use_dark else p_rgb

        p_t = tf.add_paragraph()
        p_t.text = data.get("title", "Título de la Presentación")
        p_t.font.name = font_heading
        p_t.font.size = Pt(40)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(255, 255, 255) if use_dark else p_rgb
        p_t.space_before = Pt(8)

        sub = data.get("subtitle")
        if sub:
            p_s = tf.add_paragraph()
            p_s.text = sub
            p_s.font.name = font_body
            p_s.font.size = Pt(18)
            p_s.font.color.rgb = RGBColor(226, 232, 240) if use_dark else m_rgb
            p_s.space_before = Pt(12)

        # Metadatos (Autor, Organización, Fecha)
        meta_parts = []
        if data.get("author"):
            meta_parts.append(f"Por: {data['author']}")
        if data.get("organization"):
            meta_parts.append(data["organization"])
        if data.get("date"):
            meta_parts.append(str(data["date"]))

        if meta_parts:
            meta_box = slide.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(10.5), Inches(0.8))
            tf_m = meta_box.text_frame
            p_m = tf_m.paragraphs[0]
            p_m.text = "  |  ".join(meta_parts)
            p_m.font.name = font_body
            p_m.font.size = Pt(11)
            p_m.font.color.rgb = RGBColor(203, 213, 225) if use_dark else m_rgb

    elif norm_layout == "agenda":
        _add_header(slide, data.get("title", "Agenda / Temario"), theme, data.get("subtitle"))
        items = data.get("items", [])
        start_y = 2.0
        gap = 0.95
        for idx, item in enumerate(items[:5]):
            y_pos = start_y + (idx * gap)
            # Tarjeta de fila
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), Inches(y_pos), Inches(11.7), Inches(0.8)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = bg_light_rgb
            card.line.color.rgb = border_rgb
            card.line.width = Pt(1)

            # Número
            num_box = slide.shapes.add_textbox(Inches(1.0), Inches(y_pos + 0.1), Inches(0.6), Inches(0.6))
            p_n = num_box.text_frame.paragraphs[0]
            p_n.text = f"{idx+1:02d}"
            p_n.font.name = font_heading
            p_n.font.size = Pt(16)
            p_n.font.bold = True
            p_n.font.color.rgb = s_rgb

            # Texto del ítem
            txt_box = slide.shapes.add_textbox(Inches(1.8), Inches(y_pos + 0.12), Inches(10.2), Inches(0.6))
            p_txt = txt_box.text_frame.paragraphs[0]
            if isinstance(item, dict):
                p_txt.text = item.get("title", "")
                p_txt.font.name = font_heading
                p_txt.font.size = Pt(14)
                p_txt.font.bold = True
                p_txt.font.color.rgb = txt_rgb
                desc = item.get("description")
                if desc:
                    p_d = txt_box.text_frame.add_paragraph()
                    p_d.text = desc
                    p_d.font.name = font_body
                    p_d.font.size = Pt(10.5)
                    p_d.font.color.rgb = m_rgb
            else:
                p_txt.text = str(item)
                p_txt.font.name = font_heading
                p_txt.font.size = Pt(14)
                p_txt.font.bold = True
                p_txt.font.color.rgb = txt_rgb

    elif norm_layout in ("two_columns", "three_columns"):
        _add_header(slide, data.get("title", "Análisis Comparativo"), theme, data.get("subtitle"))
        columns = data.get("columns", [])
        num_cols = 2 if norm_layout == "two_columns" else 3
        col_width = 5.6 if num_cols == 2 else 3.65
        gap = 0.5 if num_cols == 2 else 0.38
        left_start = 0.8

        for idx, col in enumerate(columns[:num_cols]):
            col_left = left_start + idx * (col_width + gap)
            # Tarjeta de columna
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(col_left), Inches(1.9), Inches(col_width), Inches(4.9)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = bg_light_rgb
            card.line.color.rgb = border_rgb
            card.line.width = Pt(1)

            # Contenido dentro de la columna
            tb = slide.shapes.add_textbox(Inches(col_left + 0.25), Inches(2.1), Inches(col_width - 0.5), Inches(4.5))
            tf = tb.text_frame
            tf.word_wrap = True

            # Título de columna
            p_t = tf.paragraphs[0]
            p_t.text = col.get("title", f"Columna {idx+1}")
            p_t.font.name = font_heading
            p_t.font.size = Pt(16)
            p_t.font.bold = True
            p_t.font.color.rgb = p_rgb

            items = col.get("items", [])
            for item in items:
                p_it = tf.add_paragraph()
                p_it.text = f"•  {item}"
                p_it.font.name = font_body
                p_it.font.size = Pt(12)
                p_it.font.color.rgb = txt_rgb
                p_it.space_before = Pt(8)

    elif norm_layout == "kpi_metrics":
        _add_header(slide, data.get("title", "Resultados y Métricas Clave"), theme, data.get("subtitle"))
        cards = data.get("cards", [])
        num_cards = min(4, max(1, len(cards)))
        card_width = (11.7 - (num_cards - 1) * 0.4) / num_cards

        for idx, card_data in enumerate(cards[:num_cards]):
            c_left = 0.8 + idx * (card_width + 0.4)
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(c_left), Inches(2.2), Inches(card_width), Inches(3.8)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = bg_light_rgb
            card.line.color.rgb = border_rgb
            card.line.width = Pt(1.5)

            tb = slide.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.5), Inches(card_width - 0.4), Inches(3.2))
            tf = tb.text_frame
            tf.word_wrap = True

            # Label
            p_lbl = tf.paragraphs[0]
            p_lbl.text = str(card_data.get("label", "")).upper()
            p_lbl.font.name = font_heading
            p_lbl.font.size = Pt(11)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = m_rgb
            p_lbl.alignment = PP_ALIGN.CENTER

            # Big Number
            p_val = tf.add_paragraph()
            p_val.text = str(card_data.get("value", ""))
            p_val.font.name = font_heading
            p_val.font.size = Pt(36)
            p_val.font.bold = True
            p_val.font.color.rgb = p_rgb
            p_val.alignment = PP_ALIGN.CENTER
            p_val.space_before = Pt(14)

            # Description
            desc = card_data.get("description")
            if desc:
                p_desc = tf.add_paragraph()
                p_desc.text = str(desc)
                p_desc.font.name = font_body
                p_desc.font.size = Pt(11)
                p_desc.font.color.rgb = s_rgb
                p_desc.alignment = PP_ALIGN.CENTER
                p_desc.space_before = Pt(12)

    elif norm_layout == "conclusion":
        _add_header(slide, data.get("title", "Conclusiones y Próximos Pasos"), theme, data.get("subtitle"))
        # Tarjeta izquierda de conclusiones
        card_l = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.9), Inches(7.5), Inches(4.8)
        )
        card_l.fill.solid()
        card_l.fill.fore_color.rgb = bg_light_rgb
        card_l.line.color.rgb = border_rgb

        tb_l = slide.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(6.9), Inches(4.2))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        p_tl = tf_l.paragraphs[0]
        p_tl.text = "Puntos Clave"
        p_tl.font.name = font_heading
        p_tl.font.size = Pt(18)
        p_tl.font.bold = True
        p_tl.font.color.rgb = p_rgb

        for point in data.get("points", ["Objetivos cumplidos satisfactoriamente", "Optimización de costos lograda", "Lanzamiento programado según cronograma"]):
            p_pt = tf_l.add_paragraph()
            p_pt.text = f"✔  {point}"
            p_pt.font.name = font_body
            p_pt.font.size = Pt(13)
            p_pt.font.color.rgb = txt_rgb
            p_pt.space_before = Pt(10)

        # Tarjeta derecha de contacto o Call To Action
        card_r = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(8.7), Inches(1.9), Inches(3.8), Inches(4.8)
        )
        card_r.fill.solid()
        card_r.fill.fore_color.rgb = p_rgb
        card_r.line.fill.background()

        tb_r = slide.shapes.add_textbox(Inches(9.0), Inches(2.4), Inches(3.2), Inches(3.8))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        p_tr = tf_r.paragraphs[0]
        p_tr.text = data.get("cta_title", "Próximo Paso")
        p_tr.font.name = font_heading
        p_tr.font.size = Pt(18)
        p_tr.font.bold = True
        p_tr.font.color.rgb = RGBColor(255, 255, 255)

        p_rc = tf_r.add_paragraph()
        p_rc.text = data.get("cta_text", "Aprobación de la propuesta para dar inicio a la Fase 1.")
        p_rc.font.name = font_body
        p_rc.font.size = Pt(12)
        p_rc.font.color.rgb = RGBColor(226, 232, 240)
        p_rc.space_before = Pt(12)

    else:
        # 'content' estándar
        _add_header(slide, data.get("title", "Contenido"), theme, data.get("subtitle"))
        body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.8))
        tf = body_box.text_frame
        tf.word_wrap = True

        bullets = data.get("bullets", [])
        if bullets:
            for idx, b in enumerate(bullets):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = f"•  {b}"
                p.font.name = font_body
                p.font.size = Pt(15)
                p.font.color.rgb = txt_rgb
                p.space_before = Pt(12)
        elif data.get("text"):
            p = tf.paragraphs[0]
            p.text = data["text"]
            p.font.name = font_body
            p.font.size = Pt(14)
            p.font.color.rgb = txt_rgb

    # Speaker notes
    if speaker_notes:
        slide.notes_slide.notes_text_frame.text = speaker_notes

    return slide


def add_table_to_slide(
    slide,
    table_data: List[List[str]],
    theme_name: str = "corporate_blue",
    left: float = 0.8,
    top: float = 2.0,
    width: float = 11.7,
    height: float = 4.5,
):
    """
    Inserta una tabla estilizada con colores de la paleta temática en la diapositiva.
    """
    if not PPTX_AVAILABLE or not table_data:
        return None

    theme = get_theme(theme_name)
    rows = len(table_data)
    cols = max(len(r) for r in table_data)

    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table

    p_rgb = RGBColor(*hex_to_rgb(theme["primary"]))
    bg_light_rgb = RGBColor(*hex_to_rgb(theme["bg_light"]))
    txt_rgb = RGBColor(*hex_to_rgb(theme["text_primary"]))
    font_name = theme["font_body"]

    for r_idx, row_vals in enumerate(table_data):
        for c_idx, val in enumerate(row_vals):
            if c_idx < cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = str(val)
                p = cell.text_frame.paragraphs[0]
                p.font.name = font_name
                
                if r_idx == 0:
                    # Cabecera
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = p_rgb
                    p.font.bold = True
                    p.font.size = Pt(12)
                    p.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    # Filas de datos con alternancia
                    cell.fill.solid()
                    if r_idx % 2 == 0:
                        cell.fill.fore_color.rgb = bg_light_rgb
                    else:
                        cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    p.font.size = Pt(11)
                    p.font.color.rgb = txt_rgb

    return table_shape
