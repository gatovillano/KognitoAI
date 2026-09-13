"""
Motor de diseño premium para hojas de cálculo Excel (.xlsx).
Proporciona tarjetas KPI superiores, gráficos nativos openpyxl (bar, column, line, pie),
formato condicional inteligente y auto-ajuste ejecutivo de columnas y paneles.
"""

import logging
from typing import Optional, List, Dict, Any, Union
import re

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter, range_boundaries, coordinate_to_tuple
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
    XLSX_AVAILABLE = True
except Exception as e:
    XLSX_AVAILABLE = False

from skills.onlyoffice_skill.design.themes import (
    get_theme,
    hex_to_rgb,
    lighten_color,
    darken_color,
)

logger = logging.getLogger(__name__)


def _parse_range_string(range_str: str) -> tuple:
    """Parsea 'A1:C10' a (min_col, min_row, max_col, max_row)."""
    # Si viene con nombre de hoja 'Hoja1!A1:C10', limpiarlo
    if "!" in range_str:
        range_str = range_str.split("!")[-1]
    return range_boundaries(range_str)


def add_kpi_cards(
    sheet,
    cards: List[Dict[str, Any]],
    theme_name: str = "corporate_blue",
    start_row: int = 2,
    start_col: int = 2,
    col_width_per_card: int = 2,
    col_gap: int = 1,
):
    """
    Inserta un panel horizontal de tarjetas KPI en las filas superiores de la hoja.
    Cada tarjeta se crea como un bloque de celdas combinadas:
      - Fila start_row: Etiqueta (8.5pt bold mayúsculas, color secundario)
      - Fila start_row+1: Valor Big Number (18pt bold, color primario)
      - Fila start_row+2: Cambio o subtítulo (9pt bold verde/rojo/accent)
    """
    if not XLSX_AVAILABLE or not cards:
        return

    theme = get_theme(theme_name)
    font_family = theme["font_heading"]

    bg_fill = PatternFill(
        start_color=theme["bg_light"],
        end_color=theme["bg_light"],
        fill_type="solid"
    )
    border_side = Side(style="thin", color=theme["border"])
    card_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    current_col = start_col

    for card in cards:
        end_col = current_col + col_width_per_card - 1
        end_row = start_row + 2

        # 1. Aplicar fondo y bordes a todo el rectángulo de la tarjeta
        for r in range(start_row, end_row + 1):
            for c in range(current_col, end_col + 1):
                cell = sheet.cell(row=r, column=c)
                cell.fill = bg_fill
                # Bordes exteriores de la tarjeta
                cell.border = Border(
                    left=border_side if c == current_col else None,
                    right=border_side if c == end_col else None,
                    top=border_side if r == start_row else None,
                    bottom=border_side if r == end_row else None,
                )

        # 2. Combinar celdas horizontalmente por fila
        sheet.merge_cells(start_row=start_row, start_column=current_col, end_row=start_row, end_column=end_col)
        sheet.merge_cells(start_row=start_row + 1, start_column=current_col, end_row=start_row + 1, end_column=end_col)
        sheet.merge_cells(start_row=start_row + 2, start_column=current_col, end_row=start_row + 2, end_column=end_col)

        # 3. Asignar contenido y estilo
        # Fila 1: Label
        c_lbl = sheet.cell(row=start_row, column=current_col)
        c_lbl.value = str(card.get("label", "")).upper()
        c_lbl.font = Font(name=font_family, size=8.5, bold=True, color=theme["text_muted"])
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")

        # Fila 2: Big Number
        c_val = sheet.cell(row=start_row + 1, column=current_col)
        c_val.value = card.get("value", "")
        c_val.font = Font(name=font_family, size=18, bold=True, color=theme["primary"])
        c_val.alignment = Alignment(horizontal="center", vertical="center")

        # Fila 3: Change / Subtext
        change = card.get("change")
        c_ch = sheet.cell(row=start_row + 2, column=current_col)
        c_ch.value = str(change) if change is not None else ""
        ch_str = str(change or "")
        if ch_str.startswith("+"):
            ch_color = "059669"  # emerald-600
        elif ch_str.startswith("-"):
            ch_color = "DC2626"  # red-600
        else:
            ch_color = theme["accent"]
        c_ch.font = Font(name=font_family, size=9, bold=True, color=ch_color)
        c_ch.alignment = Alignment(horizontal="center", vertical="center")

        current_col = end_col + 1 + col_gap


def add_chart(
    sheet,
    chart_type: str,
    data_range: str,
    categories_range: Optional[str] = None,
    target_cell: str = "F8",
    title: str = "Gráfico",
    theme_name: str = "corporate_blue",
    width: float = 16,
    height: float = 10,
):
    """
    Inserta un gráfico nativo de openpyxl (bar, column, line, pie).
    chart_type: 'bar', 'column', 'line', 'pie'
    data_range: ej. 'B1:B10' (incluyendo cabecera si hay títulos de serie)
    categories_range: ej. 'A2:A10' (etiquetas de eje x / rebanadas de torta)
    """
    if not XLSX_AVAILABLE:
        return None

    theme = get_theme(theme_name)
    norm_type = chart_type.lower().strip()

    if norm_type == "column":
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
    elif norm_type == "bar":
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
    elif norm_type == "line":
        chart = LineChart()
        chart.style = 13
    elif norm_type == "pie":
        chart = PieChart()
        chart.style = 10
    else:
        # Fallback a columnas
        chart = BarChart()
        chart.type = "col"

    chart.title = title
    chart.width = width
    chart.height = height

    # Parsear rangos
    min_col, min_row, max_col, max_row = _parse_range_string(data_range)
    data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data, titles_from_data=True)

    if categories_range:
        c_min_col, c_min_row, c_max_col, c_max_row = _parse_range_string(categories_range)
        categories = Reference(sheet, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
        chart.set_categories(categories)

    sheet.add_chart(chart, target_cell)
    return chart


def apply_conditional_formatting(
    sheet,
    range_address: str,
    rule_type: str = "status_highlights",
    options: Optional[Dict[str, Any]] = None,
):
    """
    Aplica reglas de formato condicional sobre un rango de celdas.
    rule_type:
      - 'status_highlights': Resalta estados comunes ('Completado', 'En Progreso', 'Atrasado', 'Pendiente', 'Aprobado', 'Rechazado').
      - 'color_scale': Aplica gradiente de color (ej. verde a rojo o azul a blanco).
      - 'cell_is': Regla personalizada basada en operador y valor.
    """
    if not XLSX_AVAILABLE:
        return

    options = options or {}

    if rule_type == "status_highlights":
        # Regla 1: Verde para éxitos / completados
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        green_font = Font(color="166534", bold=True)
        for val in ["Completado", "Aprobado", "Exitoso", "Alto"]:
            sheet.conditional_formatting.add(
                range_address,
                CellIsRule(operator="equal", formula=[f'"{val}"'], fill=green_fill, font=green_font)
            )

        # Regla 2: Amarillo / Ámbar para en progreso / pendientes
        amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        amber_font = Font(color="92400E", bold=True)
        for val in ["En Progreso", "Pendiente", "En Revisión", "Medio"]:
            sheet.conditional_formatting.add(
                range_address,
                CellIsRule(operator="equal", formula=[f'"{val}"'], fill=amber_fill, font=amber_font)
            )

        # Regla 3: Rojo para atrasados / rechazados
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        red_font = Font(color="991B1B", bold=True)
        for val in ["Atrasado", "Rechazado", "Crítico", "Bajo"]:
            sheet.conditional_formatting.add(
                range_address,
                CellIsRule(operator="equal", formula=[f'"{val}"'], fill=red_fill, font=red_font)
            )

    elif rule_type == "color_scale":
        # Escala de 3 colores (Rojo -> Amarillo -> Verde) o personalizada
        start_color = options.get("start_color", "F87171")
        mid_color = options.get("mid_color", "FEF08A")
        end_color = options.get("end_color", "4ADE80")

        rule = ColorScaleRule(
            start_type="min", start_color=start_color,
            mid_type="percentile", mid_value=50, mid_color=mid_color,
            end_type="max", end_color=end_color
        )
        sheet.conditional_formatting.add(range_address, rule)

    elif rule_type == "cell_is":
        operator = options.get("operator", "greaterThan")
        formula = options.get("formula", ["0"])
        fill_color = options.get("fill_color", "DCFCE7")
        font_color = options.get("font_color", "166534")

        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        font = Font(color=font_color, bold=True)
        rule = CellIsRule(operator=operator, formula=formula, fill=fill, font=font)
        sheet.conditional_formatting.add(range_address, rule)


def auto_fit_columns_and_freeze(sheet, freeze_cell: Optional[str] = None, margin: int = 3):
    """
    Ajusta automáticamente el ancho de todas las columnas utilizadas agregando margen
    para evitar textos truncados y '###' en celdas numéricas.
    Opcionalmente congela paneles en freeze_cell (ej. 'A2' o 'B5').
    """
    if not XLSX_AVAILABLE:
        return

    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                lines = val_str.split("\n")
                line_len = max(len(l) for l in lines)
                max_len = max(max_len, line_len)
            else:
                max_len = max(max_len, len(val_str))
        
        calculated_width = max(10, max_len + margin)
        # Evitar anchos absurdamente grandes
        sheet.column_dimensions[col_letter].width = min(60, calculated_width)

    if freeze_cell:
        sheet.freeze_panes = freeze_cell
