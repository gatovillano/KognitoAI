# skills/onlyoffice_skill/scripts/create_onlyoffice_document_tool.py

"""
Herramienta para crear nuevos documentos de OnlyOffice (Word, Excel, PowerPoint, Texto) desde cero
o a partir de Plantillas Maestras con diseño premium.
"""

import os
import logging
import uuid
from typing import Any, Type, Optional
from pydantic import BaseModel, Field
from langchain_core.tools import BaseTool
from core.database import SessionLocal, Document
from core.onlyoffice_storage import build_onlyoffice_relative_path, ensure_onlyoffice_account_dir, get_onlyoffice_docs_root
from skills.onlyoffice_skill.design.template_registry import (
    create_from_template,
    get_template_info,
    list_templates,
)
from skills.onlyoffice_skill.design.themes import DEFAULT_THEME, get_theme
from skills.onlyoffice_skill.design.pptx_engine import _ensure_widescreen

# Librerías de procesamiento de documentos (Office)
try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from openpyxl import Workbook
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

logger = logging.getLogger(__name__)

# Directorio para los documentos de OnlyOffice
DOCUMENTS_ROOT = str(get_onlyoffice_docs_root())
os.makedirs(DOCUMENTS_ROOT, exist_ok=True)


class CreateOnlyOfficeInput(BaseModel):
    name: str = Field(..., description="Nombre del archivo (con o sin extensión).")
    doc_type: Optional[str] = Field(
        None, 
        description="Tipo de documento a crear: 'word' (.docx), 'excel' (.xlsx), 'powerpoint' (.pptx) o 'text' (.txt). Si se pasa 'template', se deduce automáticamente."
    )
    template: Optional[str] = Field(
        None,
        description=(
            "ID de plantilla maestra para inicializar el documento con diseño premium estructurado. "
            "Opciones: 'executive_report' (Word), 'commercial_proposal' (Word), 'project_charter' (Word), "
            "'financial_budget' (Excel), 'sales_dashboard' (Excel), 'project_tracker' (Excel), "
            "'pitch_deck' (PowerPoint), 'quarterly_review' (PowerPoint)."
        )
    )
    theme: Optional[str] = Field(
        None,
        description="Tema visual premium ('corporate_blue', 'emerald_green', 'dark_slate', 'sunset_amber', 'violet_gold', 'minimalist_gray')."
    )
    title: Optional[str] = Field(
        None,
        description="Título principal para inicializar la portada o encabezado del documento."
    )
    subtitle: Optional[str] = Field(
        None,
        description="Subtítulo o descripción secundaria para la portada o presentación."
    )
    workspace_id: Optional[str] = Field(None, description="ID del workspace (opcional).")
    folder_id: Optional[str] = Field(None, description="ID de la carpeta (opcional).")


class CreateOnlyOfficeDocumentTool(BaseTool):
    name: str = "create_onlyoffice_document"
    description: str = (
        "[MÓDULO ONLYOFFICE DE KAI] Crea un nuevo documento en el módulo de documentos de la plataforma KAI. "
        "El archivo resultante aparecerá en la sección 'Documentos' de la interfaz de KAI y podrá abrirse "
        "con el editor online de OnlyOffice. Soporta: Word (.docx), Excel (.xlsx), PowerPoint (.pptx) y Texto (.txt). "
        "Permite inicializar documentos directamente desde plantillas maestras ('template') con portadas ejecutivas, "
        "tarjetas de KPI, gráficos nativos y diapositivas completas con temas visuales ('theme'). "
        "ÚSALA SOLO SI el usuario quiere un documento editable en KAI (Word/Excel/PPT). "
        "NO uses esta herramienta para: crear archivos en el disco local del usuario o servidor "
        "(usa developer_tools_skill), generar PDFs descargables (usa create_pdf_tool), "
        "ni guardar información en la memoria/base de conocimientos de KAI (usa knowledge_and_memory_skill)."
    )
    args_schema: Type[BaseModel] = CreateOnlyOfficeInput
    
    account_id: str = Field(..., description="ID de cuenta del usuario, inyectado automáticamente.")

    async def _arun(
        self, 
        name: str, 
        doc_type: Optional[str] = None, 
        template: Optional[str] = None,
        theme: Optional[str] = None,
        title: Optional[str] = None,
        subtitle: Optional[str] = None,
        workspace_id: Optional[str] = None, 
        folder_id: Optional[str] = None, 
        **kwargs: Any
    ) -> str:
        try:
            acc_id = uuid.UUID(self.account_id)
            
            # Si se pasa plantilla, validar y deducir doc_type
            if template:
                template_clean = template.strip().lower()
                tpl_info = get_template_info(template_clean)
                if not tpl_info:
                    available_tpls = ", ".join(t["id"] for t in list_templates())
                    return f"❌ Plantilla '{template}' no encontrada. Plantillas disponibles: {available_tpls}"
                if not doc_type:
                    doc_type = tpl_info["doc_type"]

            if not doc_type:
                return "❌ Debes especificar 'doc_type' ('word', 'excel', 'powerpoint', 'text') o un 'template' válido."

            # Determinar extensión y validar librerías
            doc_type = doc_type.lower().strip()
            extension = ""
            if doc_type == 'word':
                if not DOCX_AVAILABLE:
                    return "❌ Error: La librería 'python-docx' no está instalada."
                extension = 'docx'
            elif doc_type == 'excel':
                if not XLSX_AVAILABLE:
                    return "❌ Error: La librería 'openpyxl' no está instalada."
                extension = 'xlsx'
            elif doc_type == 'powerpoint':
                if not PPTX_AVAILABLE:
                    return "❌ Error: La librería 'python-pptx' no está instalada."
                extension = 'pptx'
            elif doc_type == 'text':
                extension = 'txt'
            else:
                return f"❌ Tipo de documento '{doc_type}' no soportado. Usa: 'word', 'excel', 'powerpoint' o 'text'."

            # Asegurar nombre con extensión
            if not name.lower().endswith(f".{extension}"):
                full_filename = f"{name}.{extension}"
            else:
                full_filename = name

            # Generar ruta física
            unique_filename = f"{uuid.uuid4()}.{extension}"
            user_dir = ensure_onlyoffice_account_dir(self.account_id)
            file_path = user_dir / unique_filename

            # Crear archivo físico (con plantilla o base)
            try:
                if template:
                    created = create_from_template(
                        template_id=template.strip().lower(),
                        file_path=str(file_path),
                        title=title or full_filename.rsplit(".", 1)[0],
                        subtitle=subtitle,
                        theme_name=theme or DEFAULT_THEME
                    )
                    if not created:
                        return f"❌ Error al crear documento desde la plantilla '{template}'."
                else:
                    if doc_type == 'word':
                        doc = docx.Document()
                        doc.save(file_path)
                    elif doc_type == 'excel':
                        wb = Workbook()
                        wb.save(file_path)
                    elif doc_type == 'powerpoint':
                        prs = Presentation()
                        _ensure_widescreen(prs)
                        prs.save(file_path)
                    elif doc_type == 'text':
                        with open(file_path, "w", encoding="utf-8") as f:
                            f.write("")
            except Exception as e:
                logger.error(f"Error al crear archivo físico OnlyOffice: {e}")
                return f"❌ Error al crear el archivo físico en el servidor: {str(e)}"

            # Registrar en DB
            async with SessionLocal() as db:
                new_doc = Document(
                    account_id=acc_id,
                    workspace_id=uuid.UUID(workspace_id) if workspace_id and workspace_id != "null" else None,
                    folder_id=uuid.UUID(folder_id) if folder_id and folder_id != "null" else None,
                    filename=full_filename,
                    extension=extension,
                    file_path=build_onlyoffice_relative_path(self.account_id, unique_filename)
                )
                db.add(new_doc)
                await db.commit()
                await db.refresh(new_doc)
                
                tpl_msg = f" a partir de la plantilla '{template}' (Tema: {theme or DEFAULT_THEME})" if template else ""
                return (
                    f"✅ Documento '{full_filename}' creado con éxito{tpl_msg}.\n"
                    f"ID: {new_doc.id}\n"
                    "Ahora puedes abrirlo en el editor de OnlyOffice o usar 'edit_onlyoffice_document' para ampliarlo."
                )

        except Exception as e:
            logger.error(f"Error en CreateOnlyOfficeDocumentTool: {e}", exc_info=True)
            return f"❌ Error inesperado al crear el documento: {str(e)}"

    def _run(self, *args: Any, **kwargs: Any) -> Any:
        raise NotImplementedError("Esta herramienta solo soporta ejecución asíncrona.")
