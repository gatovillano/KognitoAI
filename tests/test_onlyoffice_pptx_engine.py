import pytest
from pptx import Presentation
from pptx.util import Inches
from skills.onlyoffice_skill.design.pptx_engine import (
    add_slide,
    add_table_to_slide,
    SUPPORTED_PPTX_LAYOUTS,
)

def test_add_slide_title():
    prs = Presentation()
    data = {
        "title": "Estrategia Digital 2026",
        "subtitle": "Transformación e Innovación con IA",
        "author": "KAI Assistant",
        "organization": "Kognito AI",
        "date": "2026-09-13"
    }
    slide = add_slide(prs, layout_type="title", data=data, theme_name="corporate_blue", speaker_notes="Bienvenida al comité")
    assert slide is not None
    assert len(prs.slides) == 1
    assert slide.has_notes_slide
    assert "Bienvenida" in slide.notes_slide.notes_text_frame.text

def test_add_slide_kpi_metrics():
    prs = Presentation()
    data = {
        "title": "Métricas de Impacto",
        "cards": [
            {"label": "ROI Proyectado", "value": "340%", "description": "+45% vs benchmark"},
            {"label": "Tiempo de Adopción", "value": "14 días", "description": "-60% de reducción"},
            {"label": "Satisfacción", "value": "98.5%", "description": "CSAT Excelente"}
        ]
    }
    slide = add_slide(prs, layout_type="kpi_metrics", data=data, theme_name="emerald_green")
    assert slide is not None
    assert len(prs.slides) == 1

def test_add_slide_columns():
    prs = Presentation()
    data_2col = {
        "title": "Comparativa de Soluciones",
        "columns": [
            {"title": "Modelo Tradicional", "items": ["Costoso", "Lento", "Manual"]},
            {"title": "Modelo KAI Autónomo", "items": ["Económico", "En tiempo real", "Automatizado"]}
        ]
    }
    slide = add_slide(prs, layout_type="two_columns", data=data_2col, theme_name="dark_slate")
    assert slide is not None
    assert len(prs.slides) == 1

def test_add_table_to_slide():
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    table_data = [
        ["Fase", "Duración", "Entregable"],
        ["Fase 1", "2 semanas", "Arquitectura"],
        ["Fase 2", "4 semanas", "Despliegue"]
    ]
    table_shape = add_table_to_slide(slide, table_data, theme_name="corporate_blue")
    assert table_shape is not None
    assert len(table_shape.table.rows) == 3
    assert len(table_shape.table.columns) == 3
