import pytest
import openpyxl
from skills.onlyoffice_skill.design.xlsx_engine import (
    add_kpi_cards,
    add_chart,
    apply_conditional_formatting,
    auto_fit_columns_and_freeze,
)

def test_add_kpi_cards_excel():
    wb = openpyxl.Workbook()
    sheet = wb.active
    cards = [
        {"label": "Ventas Totales", "value": "$125,000", "change": "+14.5%"},
        {"label": "Margen Bruto", "value": "42.8%", "change": "+2.1%"},
        {"label": "Clientes Nuevos", "value": "1,420", "change": "-3.2%"},
    ]
    add_kpi_cards(sheet, cards, theme_name="corporate_blue", start_row=2, start_col=2)
    # Verificar que las celdas se hayan escrito y formateado
    val_b2 = sheet["B2"].value
    val_b3 = sheet["B3"].value
    val_b4 = sheet["B4"].value
    assert "VENTAS TOTALES" in str(val_b2).upper()
    assert val_b3 == "$125,000"
    assert val_b4 == "+14.5%"

def test_add_chart_excel():
    wb = openpyxl.Workbook()
    sheet = wb.active
    # Escribir datos de prueba
    sheet["A1"] = "Mes"
    sheet["B1"] = "Ventas"
    data = [("Ene", 100), ("Feb", 150), ("Mar", 200)]
    for row_idx, (m, v) in enumerate(data, start=2):
        sheet[f"A{row_idx}"] = m
        sheet[f"B{row_idx}"] = v

    chart = add_chart(
        sheet,
        chart_type="column",
        data_range="B1:B4",
        categories_range="A2:A4",
        target_cell="D2",
        title="Ventas por Mes",
        theme_name="corporate_blue"
    )
    assert chart is not None
    assert len(sheet._charts) == 1
    assert sheet._charts[0].title.text.to_tree().xpath(".//a:t", namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"})[0].text == "Ventas por Mes"

def test_apply_conditional_formatting_excel():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["A1"] = "Estado"
    sheet["A2"] = "Completado"
    sheet["A3"] = "En Progreso"
    sheet["A4"] = "Atrasado"

    apply_conditional_formatting(sheet, range_address="A2:A4", rule_type="status_highlights")
    assert len(sheet.conditional_formatting) > 0

def test_auto_fit_columns_and_freeze():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet["A1"] = "Encabezado muy largo para verificar ajuste"
    auto_fit_columns_and_freeze(sheet, freeze_cell="A2")
    assert sheet.column_dimensions["A"].width > 15
    assert sheet.freeze_panes == "A2"
