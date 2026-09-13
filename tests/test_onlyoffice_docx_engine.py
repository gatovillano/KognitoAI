import pytest
import docx
from skills.onlyoffice_skill.design.docx_engine import (
    add_cover_page,
    render_callout,
    add_kpi_cards,
    add_page_number_footer,
    CALLOUT_TYPES,
)

def test_add_cover_page_modern():
    doc = docx.Document()
    add_cover_page(
        doc,
        title="Informe de Gestión Estratégica",
        subtitle="Resultados Anuales y Proyecciones",
        author="KAI Assistant",
        organization="Kognito AI Corp",
        date_str="2026-09-13",
        theme_name="corporate_blue",
        style="modern"
    )
    # Debe tener párrafos y tabla de metadatos o portada
    text_content = " ".join(p.text for p in doc.paragraphs)
    assert "Informe de Gestión Estratégica" in text_content
    assert "Resultados Anuales y Proyecciones" in text_content

def test_render_callout():
    doc = docx.Document()
    table_note = render_callout(doc, "NOTE", "Esta es una nota importante sobre el proceso.", theme_name="corporate_blue")
    assert table_note is not None
    table_tip = render_callout(doc, "TIP", "Asegúrate de guardar los cambios frecuentemente.", theme_name="emerald_green")
    assert table_tip is not None
    table_warn = render_callout(doc, "WARNING", "No elimines registros sin autorización.", theme_name="sunset_amber")
    assert table_warn is not None

def test_add_kpi_cards():
    doc = docx.Document()
    cards = [
        {"value": "$4.5M", "label": "Ingresos Totales", "change": "+18.2%"},
        {"value": "99.9%", "label": "Disponibilidad", "change": "+0.4%"},
        {"value": "12.4k", "label": "Usuarios Activos", "change": "+34%"},
    ]
    table = add_kpi_cards(doc, cards, theme_name="corporate_blue")
    assert table is not None
    assert len(table.rows[0].cells) == 3

def test_add_page_number_footer():
    doc = docx.Document()
    footer = add_page_number_footer(doc, doc_title="Manual de Procedimientos", confidentiality="Confidencial")
    assert footer is not None
    footer_text = " ".join(p.text for p in footer.paragraphs)
    assert "Manual de Procedimientos" in footer_text or "Confidencial" in footer_text
